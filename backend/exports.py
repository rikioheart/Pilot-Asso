"""Phase 8 — File de priorité du Bureau, exports Excel, e-mail récapitulatif bimensuel."""
import io
import ipaddress
import os
import re
from datetime import timedelta
from html import escape
from html.parser import HTMLParser
from typing import Optional
from urllib.parse import urlparse

import httpx
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from deps import db, iso, now_utc, logger, require, require_admin, log_action, notify
from finance import CATEGORY_LABELS, period_range

router = APIRouter(prefix="/api")

EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ.get("EMERGENT_EMAIL_KEY")
EMAIL_FROM_NAME = os.environ.get("EMAIL_FROM_NAME", "Association La Voix du Chien")

HEADER_FILL = PatternFill("solid", fgColor="002060")
HEADER_FONT = Font(color="FFFFFF", bold=True)

_SHORTENERS = ("bit.ly", "tinyurl.com", "t.co", "is.gd", "cutt.ly", "goo.gl", "rebrand.ly")
_CRED_ASK = ("reply with your password", "reply with the code", "send your password", "cvv",
             "send us your password", "enter your password below", "confirm your card number",
             "your full card number", "seed phrase", "recovery phrase", "verify your card",
             "social security number", "confirm your bank details")
_HOSTISH = re.compile(r"\b(?:https?://)?((?:[a-z0-9-]+\.)+[a-z]{2,})", re.I)


def _host_ok(host: str) -> bool:
    if not host or "xn--" in host:
        return False
    try:
        ipaddress.ip_address(host)
        return False
    except ValueError:
        pass
    return not any(host == s or host.endswith("." + s) for s in _SHORTENERS)


def _same_site(shown: str, real: str) -> bool:
    return shown == real or real.endswith("." + shown) or shown.endswith("." + real)


class _EmailScan(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.urls, self.anchors = set(), [], []
        self._href, self._text = None, []

    def handle_starttag(self, tag, attrs):
        self.tags.add(tag.lower())
        self.urls += [v for k, v in attrs if k.lower() in ("href", "src") and v]
        if tag.lower() == "a":
            self._href = dict((k.lower(), v) for k, v in attrs).get("href")
            self._text = []

    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag):
        if tag.lower() == "a" and self._href is not None:
            self.anchors.append((self._href, "".join(self._text)))
            self._href, self._text = None, []


def _assert_safe_email(subject: str, html: str) -> None:
    scan = _EmailScan()
    scan.feed(html)
    if scan.tags & {"form", "input", "textarea", "select"}:
        raise ValueError("No forms or input fields in email (G2)")
    body = f"{subject}\n{html}".lower()
    for p in _CRED_ASK:
        if p in body:
            raise ValueError(f"Email asks the recipient for credentials: {p!r} (G2)")
    for url in scan.urls:
        low = url.strip().lower()
        if low.startswith(("mailto:", "tel:", "cid:", "#")):
            continue
        if not low.startswith("https://"):
            raise ValueError(f"Email links/assets must be absolute https: {url!r} (G3)")
        host = urlparse(low).hostname or ""
        if not _host_ok(host) or urlparse(low).username is not None:
            raise ValueError(f"Shortened, numeric-host or credential-bearing URL: {url!r} (G3)")
    for href, text in scan.anchors:
        real = urlparse(href.strip().lower()).hostname or ""
        if not real:
            continue
        for m in _HOSTISH.finditer(text):
            if not _same_site(m.group(1).lower(), real):
                raise ValueError(f"Anchor text {m.group(1)!r} ≠ real link host {real!r} (G3)")


# ---------------------------------------------------------------- File de priorité
@router.get("/priorities")
async def priorities(admin: dict = Depends(require_admin)):
    today = iso(now_utc())[:10]
    horizon = iso(now_utc() + timedelta(days=30))[:10]
    low_stock = 0
    async for item in db.stock_items.find({"status": {"$ne": "ARCHIVED"}}, {"_id": 0}):
        if item["quantity"] <= item.get("alert_threshold", 0):
            low_stock += 1
    groups = [
        {"key": "validations", "label": "Tâches à valider", "urgency": "HIGH", "link": "/admin/validation",
         "count": await db.tasks.count_documents({"status": "PENDING_VALIDATION"})},
        {"key": "members", "label": "Adhésions en attente", "urgency": "HIGH", "link": "/admin/members?status=PENDING",
         "count": await db.users.count_documents({"status": "PENDING"})},
        {"key": "terrain", "label": "Réservations de terrain à confirmer", "urgency": "HIGH", "link": "/terrains",
         "count": await db.terrain_reservations.count_documents({"status": "PENDING"})},
        {"key": "reimbursements", "label": "Remboursements à traiter", "urgency": "HIGH",
         "link": "/finance/reimbursements",
         "count": await db.reimbursements.count_documents({"status": {"$in": ["PENDING", "APPROVED"]}})},
        {"key": "stock", "label": "Articles sous le seuil d'alerte", "urgency": "HIGH", "link": "/stock",
         "count": low_stock},
        {"key": "overdue", "label": "Tâches en retard", "urgency": "HIGH", "link": "/tasks",
         "count": await db.tasks.count_documents(
             {"deadline": {"$lt": today, "$ne": None},
              "status": {"$nin": ["COMPLETED", "ARCHIVED", "CANCELLED"]}})},
        {"key": "articles", "label": "Articles à relire", "urgency": "NORMAL", "link": "/blog",
         "count": await db.articles.count_documents({"status": "PENDING_REVIEW"})},
        {"key": "documents", "label": "Documents à renouveler sous 30 jours", "urgency": "NORMAL",
         "link": "/documents",
         "count": await db.documents.count_documents(
             {"expiry_date": {"$ne": None, "$lte": horizon}, "status": {"$nin": ["ARCHIVE", "SIGNE"]}})},
        {"key": "help", "label": "Besoins d'aide ouverts", "urgency": "NORMAL", "link": "/admin/help",
         "count": await db.help_requests.count_documents({"status": "OPEN"})},
        {"key": "social", "label": "Publications à valider", "urgency": "NORMAL", "link": "/social",
         "count": await db.social_posts.count_documents({"status": "TO_VALIDATE"})},
        {"key": "activities", "label": "Activités proposées à arbitrer", "urgency": "NORMAL", "link": "/activities",
         "count": await db.activities.count_documents({"status": "PROPOSED"})},
        {"key": "proposals", "label": "Avantages proposés par des pros", "urgency": "NORMAL", "link": "/partners",
         "count": await db.advantage_proposals.count_documents({"status": "PENDING"})},
    ]
    items = [g for g in groups if g["count"] > 0]
    return {"items": items,
            "high_total": sum(g["count"] for g in items if g["urgency"] == "HIGH"),
            "normal_total": sum(g["count"] for g in items if g["urgency"] == "NORMAL")}


# ---------------------------------------------------------------- Exports Excel
def _sheet(workbook, title: str, headers: list):
    sheet = workbook.active if not workbook.sheetnames or workbook.active.max_row == 1 \
        and workbook.active.title == "Sheet" else workbook.create_sheet()
    sheet.title = title[:31]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    return sheet


def _finalize(workbook, filename: str) -> Response:
    for sheet in workbook.worksheets:
        for column in sheet.columns:
            width = max((len(str(c.value)) for c in column if c.value is not None), default=10)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 12), 55)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


EXPORT_KINDS = ["finance", "shares", "reimbursements", "members", "stock", "participants",
                "documents", "statistics"]


@router.get("/exports/{kind}")
async def export_excel(kind: str, year: Optional[int] = None, month: Optional[int] = None,
                       quarter: Optional[int] = None, admin: dict = Depends(require("stats.view"))):
    if kind not in EXPORT_KINDS:
        raise HTTPException(status_code=400, detail="Export inconnu")
    if admin["role"] != "ADMIN_BUREAU":
        raise HTTPException(status_code=403, detail="Réservé au Bureau")
    year = year or now_utc().year
    start, end = period_range(year, month, quarter)
    period = {"$gte": start, "$lte": end}
    workbook = Workbook()

    if kind == "finance":
        sheet = _sheet(workbook, "Écritures", ["Date", "Sens", "Montant (€)", "Catégorie",
                                               "Description", "Moyen", "Saisi par"])
        async for t in db.transactions.find({"status": {"$ne": "ARCHIVED"}, "date": period},
                                            {"_id": 0}).sort("date", 1):
            sheet.append([t.get("date"), "Recette" if t["direction"] == "IN" else "Dépense",
                          t["amount"], CATEGORY_LABELS.get(t.get("category"), t.get("category")),
                          t.get("description"), t.get("payment_method"), t.get("created_by_name")])
        recap = _sheet(workbook, "Récapitulatif", ["Mois", "Recettes (€)", "Dépenses (€)", "Net (€)"])
        totals = {}
        async for t in db.transactions.find({"status": {"$ne": "ARCHIVED"}, "date": period}, {"_id": 0}):
            row = totals.setdefault((t.get("date") or "")[:7], [0.0, 0.0])
            row[0 if t["direction"] == "IN" else 1] += t["amount"]
        for label, (inflow, outflow) in sorted(totals.items()):
            recap.append([label, round(inflow, 2), round(outflow, 2), round(inflow - outflow, 2)])

    elif kind == "shares":
        sheet = _sheet(workbook, "Parts professionnels",
                       ["Date", "Intitulé", "Professionnel", "Mode", "Valeur", "Montant (€)", "Statut"])
        async for line in db.distribution_lines.find({"date": period}, {"_id": 0}).sort("date", 1):
            sheet.append([line.get("date"), line.get("label"), line.get("professional_name"),
                          "Pourcentage" if line.get("mode") == "PERCENT" else "Montant fixe",
                          line.get("value"), line.get("computed_amount"),
                          "Réglée" if line.get("status") == "PAID" else "En attente"])

    elif kind == "reimbursements":
        sheet = _sheet(workbook, "Remboursements",
                       ["Demande", "Bénéficiaire", "Motif", "Montant (€)", "Statut", "Réglé le"])
        async for r in db.reimbursements.find({"request_date": period}, {"_id": 0}).sort("request_date", 1):
            sheet.append([r.get("request_date"), r.get("beneficiary_name"), r.get("reason"),
                          r.get("amount"), r.get("status"), r.get("paid_date")])

    elif kind == "members":
        sheet = _sheet(workbook, "Adhérents", ["Nom", "E-mail", "Rôle", "Niveau", "Statut",
                                               "Ville", "Département", "Adhésion", "Dernière connexion"])
        async for user in db.users.find({}, {"_id": 0}).sort("created_at", 1):
            profile = await db.profiles.find_one({"user_id": user["user_id"]}, {"_id": 0}) or {}
            sheet.append([profile.get("display_name"), user["email"], user["role"], user["access_level"],
                          user["status"], profile.get("city"), profile.get("department"),
                          (user.get("created_at") or "")[:10], (user.get("last_login") or "")[:10]])

    elif kind == "stock":
        sheet = _sheet(workbook, "Inventaire", ["Article", "Catégorie", "Quantité", "Unité",
                                                "Seuil", "Emplacement", "Origine"])
        async for item in db.stock_items.find({"status": {"$ne": "ARCHIVED"}}, {"_id": 0}).sort("name", 1):
            sheet.append([item["name"], item.get("category"), item["quantity"], item.get("unit"),
                          item.get("alert_threshold"), item.get("location"), item.get("origin")])
        movements = _sheet(workbook, "Mouvements", ["Date", "Article", "Type", "Quantité",
                                                   "Avant", "Après", "Motif", "Profil", "Par"])
        async for m in db.stock_movements.find({}, {"_id": 0}).sort("created_at", -1).limit(2000):
            movements.append([(m.get("created_at") or "")[:10], m.get("item_name"), m.get("direction"),
                              m.get("quantity"), m.get("quantity_before"), m.get("quantity_after"),
                              m.get("reason"), m.get("related_user_name"), m.get("created_by_name")])

    elif kind == "participants":
        sheet = _sheet(workbook, "Participants", ["Type", "Intitulé", "Date", "Participant",
                                                  "Rôle", "Inscription", "Présence"])
        async for p in db.participations.find({}, {"_id": 0}).sort("registered_at", -1).limit(3000):
            if p.get("activity_id"):
                parent = await db.activities.find_one({"activity_id": p["activity_id"]},
                                                      {"_id": 0, "title": 1, "date": 1})
                kind_label, date = "Activité", (parent or {}).get("date")
            else:
                parent = await db.events.find_one({"event_id": p.get("event_id")},
                                                  {"_id": 0, "title": 1, "start_date": 1})
                kind_label, date = "Événement", ((parent or {}).get("start_date") or "")[:10]
            profile = await db.profiles.find_one({"user_id": p["user_id"]}, {"_id": 0, "display_name": 1})
            sheet.append([kind_label, (parent or {}).get("title"), date,
                          (profile or {}).get("display_name"), p.get("role"),
                          (p.get("registered_at") or "")[:10], p.get("attendance_status")])

    elif kind == "documents":
        sheet = _sheet(workbook, "Registre", ["Titre", "Catégorie", "Date", "Statut", "Expiration", "Notes"])
        async for d in db.documents.find({}, {"_id": 0}).sort("date", -1):
            sheet.append([d["title"], d.get("category"), d.get("date"), d.get("status"),
                          d.get("expiry_date"), d.get("notes")])

    else:
        sheet = _sheet(workbook, "Réseau", ["Indicateur", "Valeur"])
        rows = [
            ("Membres", await db.users.count_documents({})),
            ("Professionnels actifs", await db.users.count_documents(
                {"role": "PROFESSIONNEL", "status": "ACTIVE"})),
            ("Particuliers actifs", await db.users.count_documents(
                {"role": "PARTICULIER", "status": "ACTIVE"})),
            ("Projets", await db.projects.count_documents({})),
            ("Tâches terminées", await db.tasks.count_documents({"status": "COMPLETED"})),
            ("Activités", await db.activities.count_documents({})),
            ("Événements", await db.events.count_documents({})),
            ("Articles publiés", await db.articles.count_documents({"status": "PUBLISHED"})),
            ("Formations", await db.formations.count_documents({})),
            ("Tampons de fidélité", await db.loyalty_stamps.count_documents({})),
            ("Partenaires actifs", await db.partners.count_documents({"status": "ACTIF"})),
            ("Avantages actifs", await db.advantages.count_documents({"status": "ACTIVE"})),
        ]
        for label, value in rows:
            sheet.append([label, value])

    await log_action(admin, "EXPORT", "exports", kind, new_value={"year": year})
    return _finalize(workbook, f"lavoixduchien-{kind}-{year}.xlsx")


# ---------------------------------------------------------------- E-mail récapitulatif
async def send_email(*, to: str, subject: str, html: str) -> Optional[str]:
    payload = {"to": [to], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            resp = await client.post(f"{EMAIL_BASE_URL}/api/v1/email/send",
                                     headers={"X-Email-Key": EMAIL_KEY}, json=payload)
        resp.raise_for_status()
        return resp.json().get("id")
    except Exception as exc:
        logger.error(f"Envoi e-mail échoué : {exc}")
        return None


async def send_recap_email():
    """Récapitulatif bimensuel envoyé à l'adresse principale de l'association."""
    recipient = os.environ["ADMIN_EMAIL"]
    since = iso(now_utc() - timedelta(days=14))
    today = iso(now_utc())[:10]
    done = await db.tasks.find({"status": "COMPLETED", "completed_at": {"$gte": since}},
                               {"_id": 0, "title": 1}).limit(30).to_list(30)
    validations = await db.tasks.count_documents({"status": "PENDING_VALIDATION"})
    pending_members = await db.users.count_documents({"status": "PENDING"})
    reservations = await db.terrain_reservations.count_documents({"status": "PENDING"})
    reimbursements = await db.reimbursements.count_documents({"status": "PENDING"})
    articles = await db.articles.count_documents({"status": "PENDING_REVIEW"})
    low_stock = len([i async for i in db.stock_items.find({"status": {"$ne": "ARCHIVED"}}, {"_id": 0})
                     if i["quantity"] <= i.get("alert_threshold", 0)])
    new_members = await db.users.count_documents({"created_at": {"$gte": since}})

    done_html = "".join(f"<li>{escape(t['title'])}</li>" for t in done) or \
        "<li>Aucune tâche terminée sur la période.</li>"
    waiting = [("tâche(s) à valider", validations), ("adhésion(s) à examiner", pending_members),
               ("réservation(s) de terrain", reservations), ("remboursement(s)", reimbursements),
               ("article(s) à relire", articles), ("article(s) de stock sous le seuil", low_stock)]
    waiting_html = "".join(f"<li><strong>{count}</strong> {label}</li>"
                           for label, count in waiting if count) or \
        "<li>Rien en attente : tout est à jour, félicitations !</li>"
    html = (
        '<table role="presentation" width="100%"><tr><td style="padding:24px;'
        'font-family:Arial,sans-serif;color:#333333">'
        f'<h1 style="color:#002060;font-size:22px;margin:0 0 8px">Récapitulatif — {escape(EMAIL_FROM_NAME)}</h1>'
        f'<p style="color:#666666;margin:0 0 20px">Période des 14 derniers jours, arrêtée au {today}.</p>'
        f'<h2 style="color:#800020;font-size:16px">Ce qui a avancé</h2><ul>{done_html}</ul>'
        f'<p style="margin:0 0 16px">{new_members} nouvelle(s) inscription(s) sur la période.</p>'
        f'<h2 style="color:#800020;font-size:16px">Ce qui attend le Bureau</h2><ul>{waiting_html}</ul>'
        '<p style="background:#f4f6fb;border-left:3px solid #800020;padding:12px 16px;margin:20px 0">'
        'Chaque petit progrès compte. Un pas après l\'autre, l\'association avance grâce à vous.</p>'
        f'<p style="font-size:12px;color:#888888">Message automatique de {escape(EMAIL_FROM_NAME)}. '
        'Nous ne demandons jamais de mot de passe ni de coordonnées bancaires par e-mail.</p>'
        '</td></tr></table>')
    subject = f"Récapitulatif de l'association — {today}"
    _assert_safe_email(subject, html)
    email_id = await send_email(to=recipient, subject=subject, html=html)
    await db.email_log.insert_one({"kind": "RECAP", "recipient": recipient, "email_id": email_id,
                                   "sent_at": iso(now_utc())})
    logger.info(f"[CRON] Récapitulatif bimensuel envoyé à {recipient} (id={email_id})")
    return {"email_id": email_id}


async def send_member_monthly_recaps():
    """Récap mensuel d'engagement : notification dans la plateforme + e-mail (si non désabonné)."""
    from loyalty import build_monthly_recap
    sent, notified, skipped = 0, 0, 0
    async for member in db.users.find({"role": "PARTICULIER", "status": "ACTIVE"},
                                      {"_id": 0, "user_id": 1, "email": 1}):
        recap = await build_monthly_recap(member["user_id"])
        period = recap["period"]["label"]
        detail = (f"{recap['gained']} tampon(s) gagné(s) en {period}, total {recap['total_points']}."
                  if recap["gained"] else
                  f"Aucun tampon en {period} — votre total reste de {recap['total_points']}.")
        if recap.get("next_reward"):
            detail += (f" Encore {recap['missing']} tampon(s) pour « "
                       f"{recap['next_reward'].get('reward') or recap['next_reward'].get('label')} ».")
        await notify(member["user_id"], type="LOYALTY_RECAP",
                     title=f"Votre récap d'engagement — {period}", message=detail,
                     level="INFO", link="/loyalty")
        notified += 1

        profile = await db.profiles.find_one({"user_id": member["user_id"]},
                                            {"_id": 0, "preferences": 1})
        opted_out = ((profile or {}).get("preferences") or {}).get("monthly_recap_email") is False
        if opted_out:
            skipped += 1
            continue
        lines = "".join(
            f"<li>{escape(s.get('activity_title') or '')} — "
            f"{escape((s.get('created_at') or '')[:10])} "
            f"({'+' if s.get('points', 0) > 0 else ''}{s.get('points', 0)})</li>"
            for s in recap["stamps"]) or "<li>Aucun tampon ce mois-ci.</li>"
        goal = ""
        if recap.get("next_reward"):
            reward = recap["next_reward"].get("reward") or recap["next_reward"].get("label")
            goal = (f'<p style="margin:0 0 16px">Prochain palier : <strong>{escape(str(reward))}</strong>'
                    f' — encore {recap["missing"]} tampon(s).</p>')
        html = (
            '<table role="presentation" width="100%"><tr><td style="padding:24px;'
            'font-family:Arial,sans-serif;color:#333333">'
            f'<h1 style="color:#002060;font-size:22px;margin:0 0 8px">Votre engagement — {escape(period)}</h1>'
            f'<p style="color:#666666;margin:0 0 20px">Bonjour {escape(recap["display_name"] or "")}, '
            f'voici le résumé de votre carte d\'engagement.</p>'
            f'<p style="margin:0 0 16px">Total de tampons : <strong>{recap["total_points"]}</strong> '
            f'(dont {recap["gained"]} ce mois-ci).</p>'
            f'<h2 style="color:#800020;font-size:16px">Vos tampons du mois</h2><ul>{lines}</ul>'
            f'{goal}'
            '<p style="background:#f4f6fb;border-left:3px solid #800020;padding:12px 16px;margin:20px 0">'
            'Merci pour votre présence auprès de l\'association et de nos chiens.</p>'
            f'<p style="font-size:12px;color:#888888">Message automatique de {escape(EMAIL_FROM_NAME)}. '
            'Vous pouvez désactiver cet e-mail dans « Mon espace &amp; préférences ». '
            'Nous ne demandons jamais de mot de passe ni de coordonnées bancaires par e-mail.</p>'
            '</td></tr></table>')
        subject = f"Votre récap d'engagement — {period}"
        _assert_safe_email(subject, html)
        email_id = await send_email(to=member["email"], subject=subject, html=html)
        await db.email_log.insert_one({"kind": "MEMBER_RECAP", "recipient": member["email"],
                                       "email_id": email_id, "sent_at": iso(now_utc())})
        if email_id:
            sent += 1
    logger.info(f"[CRON] Récaps mensuels : {notified} notification(s), {sent} e-mail(s), "
                f"{skipped} désabonné(s)")
    return {"notified": notified, "emails": sent, "opted_out": skipped}


@router.post("/exports/test-recap")
async def test_recap(admin: dict = Depends(require_admin)):
    return await send_recap_email()


@router.post("/exports/member-recaps")
async def trigger_member_recaps(admin: dict = Depends(require_admin)):
    """Déclenchement manuel par le Bureau du récap mensuel des adhérents."""
    result = await send_member_monthly_recaps()
    await log_action(admin, "EXPORT", "exports", "member-recaps", new_value=result)
    return result
